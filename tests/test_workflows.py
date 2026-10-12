import json
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from pydantic_ai import Agent, NativeOutput
from pydantic_ai.models.test import TestModel
from app import db, documents, preextract, embeddings, granola
from app.provider import ProcessingError
from app.evaluation import EvalDataset, run_dataset, retrieval_metrics
from tests.test_api import client


DRAFT = dict(name="Alex Example", skills=[], current_job_title="Pilot", past_job_titles=[], years_experience=5,
             location_city="Austin", location_region="TX")


def pdf_bytes(text=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO(); writer.write(output); return output.getvalue()


@pytest.mark.parametrize("name,content", [("a.txt", b"Available weekends"), ("a.md", b"# Resume"), ("a.PDF", pdf_bytes("Pilot in Austin"))])
def test_extract_formats_without_saving(client, name, content):
    before = db.state_token()
    response = client.post("/documents/extract", files={"file": (name, content)})
    assert response.status_code == 200
    assert response.json()["text"].strip()
    assert response.json()["extraction_ms"] >= 0
    assert db.state_token() == before


def test_docx_paragraphs_and_tables(client):
    document = Document(); document.add_paragraph("Pilot resume")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Austin Texas"
    output = BytesIO(); document.save(output)
    response = client.post("/documents/extract", files={"file": ("resume.docx", output.getvalue())})
    assert response.status_code == 200
    assert "Pilot resume" in response.json()["text"] and "Austin Texas" in response.json()["text"]


@pytest.mark.parametrize("name,content", [("x.exe", b"test"), ("x.txt", b""), ("x.txt", b"\x00abc"), ("x.txt", b"\xff"),
    ("x.pdf", b"not a pdf"), ("x.pdf", b"%PDF-corrupt"), ("scan.pdf", pdf_bytes()), ("x.docx", b"bad"), ("x.txt", b" " )])
def test_invalid_uploads(client, name, content):
    assert client.post("/documents/extract", files={"file": (name, content)}).status_code == 422


def test_upload_limits(client, monkeypatch):
    monkeypatch.setattr(documents, "MAX_BYTES", 10)
    assert client.post("/documents/extract", files={"file": ("x.txt", b"x"*11)}).status_code == 413
    monkeypatch.setattr(documents, "MAX_CHARS", 2)
    assert client.post("/documents/extract", files={"file": ("x.txt", b"abc")}).status_code == 422


def test_docx_disguised_zip(client):
    output=BytesIO()
    with ZipFile(output, "w") as archive: archive.writestr("other.txt", "hello")
    assert client.post("/documents/extract", files={"file": ("x.docx", output.getvalue())}).status_code == 422


def test_pydantic_ai_draft_and_reviewed_save(client, monkeypatch):
    agent = Agent(TestModel(custom_output_text=json.dumps(DRAFT), profile={"supports_json_schema_output": True}), output_type=NativeOutput(preextract.CandidateDraft))
    monkeypatch.setattr(preextract, "make_agent", lambda: agent)
    before = db.state_token()
    response = client.post("/candidates/pre-extract", json={"text": "Alex Example resume"})
    assert response.status_code == 200
    draft = response.json()["draft"]
    assert draft["location_region"] == "Texas"
    assert db.state_token() == before
    monkeypatch.setattr("app.extraction.extract_fields", lambda _: pytest.fail("Reviewed fields must not be extracted again"))
    name = draft.pop("name"); draft["location_city"] = "Dallas"
    response = client.post("/candidates", json={"name": name, "resume": "Alex Example resume", "fields": draft})
    assert response.status_code == 201
    assert response.json()["location_city"] == "Dallas"
    assert client.post("/candidates", json={"name": name, "resume": "Alex Example resume", "fields": draft}).status_code == 409


def test_preextract_errors_do_not_save(client, monkeypatch):
    before = db.state_token()
    assert client.post("/candidates/pre-extract", json={"text":" "}).status_code == 422
    def fail(_): raise ProcessingError("unavailable")
    monkeypatch.setattr(preextract, "extract_draft", fail)
    assert client.post("/candidates/pre-extract", json={"text":"resume"}).status_code == 503
    assert db.state_token() == before


def test_granola_import_search_and_replay(client, monkeypatch):
    meetings = client.get("/integrations/granola/meetings").json()
    assert len(meetings) == 3 and meetings[0]["id"] == "demo-001"
    meeting = meetings[0]
    url = f"/integrations/granola/meetings/{meeting['id']}/import"
    body = {"candidate_id": 1, "note": meeting["text"]}
    response = client.post(url, json=body)
    assert response.status_code == 200 and response.json()["inserted"]
    saved = response.json()["note"]
    assert meeting["id"] in saved["note"] and saved["timestamp"] == meeting["timestamp"]
    search = client.post("/search", json={"query": saved["note"], "top_k":1}).json()
    assert search["results"][0]["id"] == 1
    assert search["results"][0]["match_source"] == "note"
    monkeypatch.setattr(embeddings, "embed_documents", lambda _: pytest.fail("Replay must not call provider"))
    repeated = client.post(url, json={**body,"note":"Edited on replay"}).json()
    assert not repeated["inserted"] and repeated["note"] == saved


def test_granola_invalid_and_failure_atomic(client, monkeypatch):
    before = db.state_token()
    url = "/integrations/granola/meetings/demo-001/import"
    assert client.post(url, json={"candidate_id":999,"note":"hello"}).status_code == 404
    assert client.post(url, json={"candidate_id":1,"note":" "}).status_code == 422
    assert client.post(url.replace('demo-001','missing'), json={"candidate_id":1,"note":"hello"}).status_code == 404
    def fail(_): raise RuntimeError("provider unavailable")
    monkeypatch.setattr(embeddings, "embed_documents", fail)
    assert client.post(url,json={"candidate_id":1,"note":"hello"}).status_code == 503
    assert db.state_token() == before
    assert db.get_meeting_note("demo-001",1) is None


def test_rebuild_clears_meeting_associations(client):
    client.post("/integrations/granola/meetings/demo-001/import",json={"candidate_id":1,"note":"hello"})
    from app.ingest import build_database
    build_database(reset=True)
    assert db.get_meeting_note("demo-001",1) is None


def test_eval_metrics():
    metrics = retrieval_metrics([4,2,3], [2,3], 3)
    assert metrics["recall_at_k"] == 1 and metrics["precision_at_k"] == pytest.approx(2/3)
    assert metrics["mrr"] == .5 and 0 < metrics["ndcg_at_k"] < 1
    assert retrieval_metrics([], [], 5) == {"empty_result_correct":1}
    assert retrieval_metrics([1], [], 5) == {"empty_result_correct":0}


def test_eval_records_failures_and_field_matches():
    dataset = EvalDataset.model_validate({"name":"test","description":"test","cases":[
        {"id":"one","kind":"retrieval","search":{"query":"pilot"},"relevant_ids":[1],"rationale":"test"},
        {"id":"two","kind":"extraction","text":"resume","expected_fields":{"name":"Alex Example"},"rationale":"test"}]})
    def fail(_): raise RuntimeError("secret provider text")
    report = run_dataset(dataset, fail, lambda _: preextract.CandidateDraft(**DRAFT))
    assert report["summary"]["errors"] == 1 and report["summary"]["field_accuracy"] == 1
    assert "secret" not in str(report)
    assert report["cost"] is None


def test_reviewed_fields_reject_typo_before_provider(client, monkeypatch):
    monkeypatch.setattr(embeddings, "embed_documents", lambda _: pytest.fail("Invalid fields called provider"))
    fields = {key:value for key,value in DRAFT.items() if key != "name"}
    fields["location_cty"] = "typo"
    assert client.post("/candidates", json={"name":"New person","resume":"resume","fields":fields}).status_code == 422


def test_meeting_db_race_returns_original_note(client):
    from tests.test_api import fake_embed
    row = {"candidate_id":1,"note":"meeting evidence","timestamp":"2026-10-12T00:00:00Z"}
    vector = fake_embed([row["note"]])
    first, inserted = db.insert_meeting_note("race-test", row, vector)
    second, repeated = db.insert_meeting_note("race-test", {**row,"note":"changed"}, vector)
    assert inserted and not repeated and first == second


def test_meeting_link_failure_rolls_back_note(client):
    from tests.test_api import fake_embed
    with db.connect() as conn:
        conn.execute("CREATE TRIGGER fail_meeting BEFORE INSERT ON meeting_imports BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    before = db.state_token()
    response = client.post("/integrations/granola/meetings/demo-001/import",json={"candidate_id":1,"note":"hello"})
    assert response.status_code == 409
    assert db.state_token() == before


@pytest.mark.parametrize("case", [
    {"id":"a","kind":"retrieval","relevant_ids":[1],"rationale":"test"},
    {"id":"a","kind":"retrieval","search":{"query":"test"},"relevant_ids":[1,1],"rationale":"test"},
    {"id":"a","kind":"extraction","text":"test","expected_fields":{"typo":"test"},"rationale":"test"},
])
def test_invalid_eval_cases(case):
    with pytest.raises(ValueError):
        EvalDataset.model_validate({"name":"invalid","description":"test","cases":[case]})
